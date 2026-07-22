# Copyright (c) 2025 Bytedance Ltd. and/or its affiliates
# SPDX-License-Identifier: MIT

"""附件保存/解析/读取模块的单元测试."""

import base64
import io

import pytest

# 1x1 红色像素 PNG
_PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGP4z8DwHwAFAAH/"
    "q842iQAAAABJRU5ErkJggg=="
)

def _make_pdf() -> bytes:
    """动态拼一个 xref 偏移正确的最小含文本 PDF（内容 "Hello PDF"）."""
    stream = b"BT /F1 24 Tf 72 720 Td (Hello PDF) Tj ET"
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
        b"/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n%s\nendstream" % (len(stream), stream),
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (i, body)
    xref_pos = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += (
        b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF"
        % (len(objs) + 1, xref_pos)
    )
    return out


_PDF_BYTES = _make_pdf()


@pytest.fixture()
def files_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("PLATFORM_FILES_DIR", str(tmp_path))
    return tmp_path


def _save(user_id, name, data):
    from src.skills.attachments import save_attachment

    return save_attachment(user_id, name, data)


class TestSaveAndParse:
    def test_txt_roundtrip(self, files_dir):
        from src.skills.attachments import load_parsed_texts

        meta = _save(1, "笔记.txt", "第一行\n第二行".encode())
        assert meta["kind"] == "document"
        assert meta["name"] == "笔记.txt"
        assert meta["chars"] > 0
        assert not meta.get("error")
        text = load_parsed_texts(1, [meta["id"]])
        assert "第二行" in text
        assert "笔记.txt" in text  # 注入时标注来源文件名

    def test_csv_and_md(self, files_dir):
        m1 = _save(1, "data.csv", b"a,b\n1,2")
        m2 = _save(1, "doc.md", b"# title\nbody")
        assert m1["chars"] and m2["chars"]

    def test_docx_parse(self, files_dir):
        import docx

        buf = io.BytesIO()
        d = docx.Document()
        d.add_paragraph("docx 段落内容")
        d.save(buf)
        meta = _save(1, "report.docx", buf.getvalue())
        from src.skills.attachments import load_parsed_texts

        assert "docx 段落内容" in load_parsed_texts(1, [meta["id"]])

    def test_xlsx_parse(self, files_dir):
        import openpyxl

        wb = openpyxl.Workbook()
        wb.active.append(["表头甲", "表头乙"])
        wb.active.append(["值1", "值2"])
        buf = io.BytesIO()
        wb.save(buf)
        meta = _save(1, "table.xlsx", buf.getvalue())
        from src.skills.attachments import load_parsed_texts

        text = load_parsed_texts(1, [meta["id"]])
        assert "表头甲" in text and "值2" in text

    def test_pdf_parse(self, files_dir):
        meta = _save(1, "paper.pdf", _PDF_BYTES)
        from src.skills.attachments import load_parsed_texts

        assert "Hello PDF" in load_parsed_texts(1, [meta["id"]])

    def test_image_saved_not_parsed(self, files_dir):
        from src.skills.attachments import load_image_data_urls

        meta = _save(1, "shot.png", _PNG_BYTES)
        assert meta["kind"] == "image"
        assert meta["chars"] == 0
        urls = load_image_data_urls(1, [meta["id"]])
        assert urls[0].startswith("data:image/png;base64,")

    def test_truncates_long_text(self, files_dir):
        from src.skills.attachments import MAX_PARSED_CHARS

        meta = _save(1, "big.txt", ("字" * (MAX_PARSED_CHARS + 5000)).encode())
        assert meta["chars"] == MAX_PARSED_CHARS

    def test_corrupt_document_returns_error(self, files_dir):
        meta = _save(1, "broken.docx", b"not a real docx")
        assert meta["kind"] == "document"
        assert meta["error"]
        assert meta["chars"] == 0


class TestValidation:
    def test_oversize_rejected(self, files_dir):
        from src.skills.attachments import MAX_FILE_BYTES

        with pytest.raises(ValueError):
            _save(1, "huge.txt", b"x" * (MAX_FILE_BYTES + 1))

    def test_unsupported_ext_rejected(self, files_dir):
        with pytest.raises(ValueError):
            _save(1, "evil.exe", b"MZ")

    def test_bad_id_rejected(self, files_dir):
        from src.skills.attachments import load_parsed_texts, split_ids_by_kind

        for bad in ["../../etc/passwd", "a/b.txt", "ABC.txt", "x.txt\n"]:
            with pytest.raises(ValueError):
                load_parsed_texts(1, [bad])
            with pytest.raises(ValueError):
                split_ids_by_kind(1, [bad])

    def test_missing_file_rejected(self, files_dir):
        from src.skills.attachments import split_ids_by_kind

        with pytest.raises(ValueError):
            split_ids_by_kind(1, ["0123456789ab-dead-beef.txt"])

    def test_user_isolation(self, files_dir):
        from src.skills.attachments import split_ids_by_kind

        meta = _save(1, "mine.txt", b"secret")
        with pytest.raises(ValueError):
            split_ids_by_kind(2, [meta["id"]])


class TestInjection:
    def test_split_by_kind(self, files_dir):
        from src.skills.attachments import split_ids_by_kind

        doc = _save(1, "a.txt", b"text")
        img = _save(1, "b.png", _PNG_BYTES)
        docs, imgs = split_ids_by_kind(1, [doc["id"], img["id"]])
        assert docs == [doc["id"]] and imgs == [img["id"]]

    def test_total_injection_cap(self, files_dir):
        from src.skills.attachments import (
            MAX_INJECT_CHARS,
            MAX_PARSED_CHARS,
            load_parsed_texts,
        )

        ids = [
            _save(1, f"f{i}.txt", ("长" * MAX_PARSED_CHARS).encode())["id"]
            for i in range(3)
        ]
        text = load_parsed_texts(1, ids)
        assert len(text) <= MAX_INJECT_CHARS + 200  # 允许标注头的少量溢出

    def test_reference_block_wrapping(self, files_dir):
        from src.skills.attachments import build_reference_block

        assert build_reference_block("") == ""
        block = build_reference_block("内容")
        assert block.startswith("\n\n【参考资料】")
        assert "内容" in block

    def test_empty_ids(self, files_dir):
        from src.skills.attachments import load_parsed_texts, split_ids_by_kind

        assert load_parsed_texts(1, []) == ""
        assert split_ids_by_kind(1, []) == ([], [])
